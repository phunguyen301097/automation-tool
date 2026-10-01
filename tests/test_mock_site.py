"""End-to-end tests of the runner against a local mock of the dashboard."""
from __future__ import annotations

import csv
import io
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qsl, urlparse

import openpyxl
import pytest
import yaml

from weplan_export.config import load_config, load_scenarios
from weplan_export.runner import run_scenarios

MOCK_HTML = (Path(__file__).parent / "mock_site" / "index.html").read_bytes()


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def do_GET(self):
        url = urlparse(self.path)
        if url.path == "/switch-country":
            # Like the dashboard's country switch: an endpoint answering 204 (the navigation
            # to it is aborted: net::ERR_ABORTED), then the page reloads itself.
            time.sleep(0.3)
            self.send_response(204)
            self.end_headers()
            return
        if url.path == "/suspended":
            body = (b"<html><body><h2>Account suspended</h2><p>Your account has been temporarily "
                    b"suspended due to a violation of the platform's terms of use.</p></body></html>")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if url.path == "/export":
            q = dict(parse_qsl(url.query))
            fmt = q.pop("format")
            rows = [["key", "value"]] + [[k, v] for k, v in q.items()]
            if fmt == "xlsx":
                wb = openpyxl.Workbook()
                for r in rows:
                    wb.active.append(r)
                buf = io.BytesIO()
                wb.save(buf)
                body, ctype = buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            else:
                buf = io.StringIO()
                csv.writer(buf).writerows(rows)
                body, ctype = buf.getvalue().encode(), "text/csv"
            self.send_response(200)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Disposition", f'attachment; filename="table_export.{fmt}"')
        else:
            body = MOCK_HTML
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


@pytest.fixture(scope="module")
def server():
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{httpd.server_port}"
    httpd.shutdown()


@pytest.fixture
def config(server, tmp_path):
    cfg = load_config(None)
    cfg["base_url"] = server
    cfg["output_dir"] = str(tmp_path / "out")
    cfg["auth"]["required"] = False
    cfg["timeouts"].update(default=10_000, data_load=20_000, download=20_000)
    return cfg


def _run(tmp_path, config, doc):
    f = tmp_path / "s.yaml"
    f.write_text(yaml.safe_dump(doc, allow_unicode=True), encoding="utf-8")
    return run_scenarios(load_scenarios([f]), config)


def _xlsx_dict(path):
    wb = openpyxl.load_workbook(path)
    return {r[0]: r[1] for r in wb.active.iter_rows(min_row=2, values_only=True)}


def test_full_flow_xlsx_with_matrix(tmp_path, config):
    doc = {
        "vars": {"from": "max-30d", "to": "max"},
        "scenarios": [{
            "name": "cov_${country}_${network}",
            "matrix": {"country": ["Cambodia"], "network": ["METFONE", "SMART"]},
            "steps": [
                {"goto": "/app/bi/coverage"},
                {"open_menu": ["Latency", "Latency Mobile (Cellular)"]},
                {"select_country": "${country}"},
                {"set_date": {"from": "${from}", "to": "${to}"}},
                {"select_filter": {"id": "carrier_filter", "options": ["${network}"]}},
                {"filters": {"coverage_filter": ["4G"], "origin_filter": ["Weplan"]}},
                {"choose_view": "macro"},
                {"wait_for_table": {}},
                {"download_table": {"format": "xlsx", "filename": "lat/${country}_${network}_${date_from}_${date_to}"}},
            ],
        }],
    }
    results = _run(tmp_path, config, doc)
    assert [r.ok for r in results] == [True, True], [r.error for r in results]
    out = Path(config["output_dir"])
    for net in ("METFONE", "SMART"):
        f = out / "lat" / f"Cambodia_{net}_2026-08-24_2026-09-23.xlsx"
        assert f.exists()
        data = _xlsx_dict(f)
        assert data["country"] == "kh"
        assert data["kpi"] == "/app/kh/latencyMobile"  # the country is part of the URL
        assert data["date"] == "2026-08-24..2026-09-23"
        assert data["carrier"] == net
        assert data["coverage"] == "4G"
        assert data["origin"] == "Weplan"
        assert data["view"] == "byCountry"


def test_csv_and_requery_after_filter_change(tmp_path, config):
    doc = {"scenarios": [{
        "name": "csv_requery",
        "steps": [
            {"goto": "/app/bi/coverage"},
            {"set_date": {"from": "2026-08-01", "to": "2026-08-31"}},
            {"choose_view": {"text": "By provinces"}},
            {"wait_for_table": {}},
            {"select_filter": {"id": "group_selector", "options": ["Cellular network"]}},
            {"apply": {}},
            {"wait_for_table": {}},
            {"download_table": {"format": "csv"}},
        ],
    }]}
    results = _run(tmp_path, config, doc)
    assert results[0].ok, results[0].error
    f = Path(results[0].downloads[0]["file"])
    assert f.suffix == ".csv"
    content = dict(csv.reader(f.read_text().splitlines()[1:]))
    assert content["group"] == "carrier"
    assert content["view"] == "byRegions"
    assert content["date"] == "2026-08-01..2026-08-31"


def test_failure_is_reported_with_artifacts(tmp_path, config):
    doc = {"scenarios": [{
        "name": "bad_network",
        "steps": [
            {"goto": "/app/bi/coverage"},
            {"select_filter": {"id": "carrier_filter", "options": ["NOPE"], "wait_options_ms": 1500}},
        ],
    }]}
    results = _run(tmp_path, config, doc)
    r = results[0]
    assert not r.ok
    assert "NOPE" in r.error and "ECONET" in r.error
    assert r.failed_step == "2. select_filter"
    assert any(a.endswith(".png") for a in r.artifacts)


def test_dashboard_error_fails_fast(tmp_path, config):
    doc = {"scenarios": [{
        "name": "no_date",
        "steps": [{"goto": "/app/bi/coverage"}, {"choose_view": "macro"}, {"wait_for_table": {}}],
    }]}
    r = _run(tmp_path, config, doc)[0]
    assert not r.ok
    assert "No date selected" in r.error


def test_date_widget_without_input_uses_popup_inputs(tmp_path, config):
    doc = {"scenarios": [{
        "name": "popup_date",
        "steps": [
            {"goto": "/app/bi/coverage?dp=popup"},
            {"set_date": {"from": "2026-08-01", "to": "2026-08-31"}},
            {"choose_view": "macro"},
            {"wait_for_table": {}},
            {"download_table": {"format": "csv"}},
        ],
    }]}
    r = _run(tmp_path, config, doc)[0]
    assert r.ok, r.error
    content = dict(csv.reader(Path(r.downloads[0]["file"]).read_text().splitlines()[1:]))
    assert content["date"] == "2026-08-01..2026-08-31"


def test_unknown_date_widget_dumps_html_and_skip_mode_works(tmp_path, config):
    doc = {"scenarios": [
        {"name": "unknown_date", "steps": [
            {"goto": "/app/bi/coverage?dp=none"},
            {"set_date": {"from": "2026-08-01", "to": "2026-08-31"}},
        ]},
    ]}
    r = _run(tmp_path, config, doc)[0]
    assert not r.ok
    dump = Path(config["output_dir"]) / "_debug" / "datepicker_unknown_date.html"
    assert str(dump) in r.error
    assert "reportrange-text" in dump.read_text()


def test_login_style_manual_download_is_kept(tmp_path, config, server):
    from playwright.sync_api import sync_playwright
    from weplan_export.runner import _keep_manual_downloads, launch_browser, new_context

    with sync_playwright() as pw:
        browser = launch_browser(pw, config)
        ctx = new_context(browser, config, use_state=False)
        _keep_manual_downloads(ctx, config)
        page = ctx.new_page()
        with page.expect_download() as dl:
            page.goto(server + "/app/bi/coverage")
            page.evaluate("u => { const a = document.createElement('a'); a.href = u; document.body.appendChild(a); a.click(); }",
                          server + "/export?format=csv&x=1")
        dl.value.path()  # wait until finished
        page.wait_for_timeout(500)
        browser.close()
    assert (Path(config["output_dir"]) / "manual" / "table_export.csv").exists()


def test_date_text_input_variant(tmp_path, config):
    doc = {"scenarios": [{
        "name": "input_date",
        "steps": [
            {"goto": "/app/bi/coverage?dp=input"},
            {"set_date": {"from": "2026-07-15", "to": "2026-08-02"}},
            {"choose_view": "macro"},
            {"wait_for_table": {}},
            {"download_table": {"format": "json"}},
        ],
    }]}
    r = _run(tmp_path, config, doc)[0]
    assert r.ok, r.error
    assert r.downloads[0]["file"].endswith(".json")
    assert "date,2026-07-15..2026-08-02" in Path(r.downloads[0]["file"]).read_text()


def test_calendar_day_after_max_date_fails_clearly(tmp_path, config):
    doc = {"scenarios": [{
        "name": "future_date",
        "steps": [
            {"goto": "/app/bi/coverage"},
            {"set_date": {"from": "2026-09-01", "to": "2026-09-25"}},
        ],
    }]}
    r = _run(tmp_path, config, doc)[0]
    assert not r.ok
    assert "2026-09-25 is not selectable" in r.error


def test_video_like_calendar_navigates_months(tmp_path, config):
    """Default mock widget reproduces the recorded dashboard picker (no known class names)."""
    doc = {"scenarios": [{
        "name": "nav_months",
        "steps": [
            {"goto": "/app/bi/coverage"},
            {"set_date": {"from": "2026-05-31", "to": "2026-07-01"}},
            {"choose_view": "macro"},
            {"wait_for_table": {}},
            {"download_table": {"format": "csv"}},
        ],
    }]}
    r = _run(tmp_path, config, doc)[0]
    assert r.ok, r.error
    assert "date,2026-05-31..2026-07-01" in Path(r.downloads[0]["file"]).read_text()


def test_daterangepicker_variant(tmp_path, config):
    doc = {"scenarios": [{
        "name": "drp",
        "steps": [
            {"goto": "/app/bi/coverage?dp=drp"},
            {"set_date": {"from": "2026-08-03", "to": "2026-09-14"}},
            {"choose_view": "macro"},
            {"wait_for_table": {}},
            {"download_table": {"format": "csv"}},
        ],
    }]}
    r = _run(tmp_path, config, doc)[0]
    assert r.ok, r.error
    assert "date,2026-08-03..2026-09-14" in Path(r.downloads[0]["file"]).read_text()


SESSION_MARK = "() => sessionStorage.setItem('mark', '1')"
SESSION_CHECK = "() => { if (sessionStorage.getItem('mark') !== '1') throw new Error('not the same page'); }"


def _two_scenarios():
    return {"scenarios": [
        {"name": "first", "steps": [{"goto": "/app/bi/coverage"}, {"js": SESSION_MARK}]},
        {"name": "second", "steps": [{"open_menu": ["Coverage time"]}, {"js": SESSION_CHECK}]},
    ]}


def test_scenarios_share_one_page_by_default(tmp_path, config):
    results = _run(tmp_path, config, _two_scenarios())
    assert [r.status for r in results] == ["PASS", "PASS"], [r.error for r in results]


def test_isolated_mode_uses_fresh_page(tmp_path, config):
    config["browser"]["isolated"] = True
    results = _run(tmp_path, config, _two_scenarios())
    assert [r.status for r in results] == ["PASS", "FAIL"]
    assert "not the same page" in results[1].error


def test_closing_the_browser_window_stops_the_run(tmp_path, config, monkeypatch):
    from weplan_export import actions
    monkeypatch.setitem(actions.STEPS, "user_closes_window", lambda ctx, args: ctx.page.close())
    doc = {"scenarios": [
        {"name": "a", "steps": [{"goto": "/app/bi/coverage"}]},
        {"name": "b", "steps": [{"goto": "/app/bi/coverage"}, {"user_closes_window": None}, {"wait": 100}]},
        {"name": "c", "steps": [{"goto": "/app/bi/coverage"}]},
        {"name": "d", "steps": [{"goto": "/app/bi/coverage"}]},
    ]}
    results = _run(tmp_path, config, doc)
    assert [r.status for r in results] == ["PASS", "STOPPED", "NOT RUN", "NOT RUN"]
    assert "window was closed" in results[1].error


def test_ctrl_c_stops_the_run(tmp_path, config, monkeypatch):
    from weplan_export import actions

    def ctrl_c(ctx, args):
        raise KeyboardInterrupt

    monkeypatch.setitem(actions.STEPS, "ctrl_c", ctrl_c)
    doc = {"scenarios": [
        {"name": "a", "steps": [{"goto": "/app/bi/coverage"}, {"ctrl_c": None}]},
        {"name": "b", "steps": [{"goto": "/app/bi/coverage"}]},
    ]}
    results = _run(tmp_path, config, doc)
    assert [r.status for r in results] == ["STOPPED", "NOT RUN"]


def _announce_flow(announce):
    return {"scenarios": [
        {"name": "with_popup", "steps": [
            {"goto": f"/app/bi/coverage?announce={announce}"},
            {"wait": 1800},  # popup is now covering the page
            {"set_date": {"from": "2026-08-03", "to": "2026-09-14"}},
            {"select_filter": {"id": "carrier_filter", "options": ["LUMITEL"]}},
            {"choose_view": "macro"},
            {"wait_for_table": {}},
            {"download_table": {"format": "csv"}},
        ]},
        {"name": "next_page_load", "steps": [
            {"open_menu": ["Coverage time"]},
            {"wait": 1800},
            {"choose_view": "macro"},
        ]},
    ]}


def test_announcement_popup_is_closed_automatically(tmp_path, config, capsys):
    results = _run(tmp_path, config, _announce_flow("1"))
    assert [r.status for r in results] == ["PASS", "PASS"], [r.error for r in results]
    out = capsys.readouterr().out
    assert "closed popup 'New Delta Analysis in Map View" in out
    # Closing it with its own button marks it as seen: it does not come back on the next page load.
    assert out.count("closed popup") == 1


def test_announcement_popup_without_working_close_button_is_removed(tmp_path, config, capsys):
    results = _run(tmp_path, config, _announce_flow("stuck"))
    assert [r.status for r in results] == ["PASS", "PASS"], [r.error for r in results]
    assert "(removed)" in capsys.readouterr().out


# file -> (page path, [(name part after VTB_<year>_T<month>_, expected mock state)])
_COV = {"All": "5G_SA|5G_NSA_CONNECTED|5G_NSA_NOT_RESTRICTED|5G_NSA_RESTRICTED|4G|3G|2G",
        "5G": "5G_SA|5G_NSA_CONNECTED|5G_NSA_NOT_RESTRICTED|5G_NSA_RESTRICTED", "4G": "4G"}


def _coverage_kpi(kpi):
    return [(f"{kpi}_{lvl}_{t}", {"coverage": _COV[t]}) for lvl in ("Net", "Province") for t in ("All", "5G", "4G")]


KPI_SCENARIOS = {
    "coverage_time.yaml": ("/app/bi/coverage", _coverage_kpi("Coverage time")),
    "signal_strength.yaml": ("/app/bi/signal", _coverage_kpi("Signal strength")),
    "data_traffic.yaml": ("/app/bi/traffic", _coverage_kpi("Data traffic")),
    "latency.yaml": ("/app/bi/latencyMobile", _coverage_kpi("Latency")),
    "packet_loss.yaml": ("/app/bi/latencyPacketLossMobile", _coverage_kpi("Packet Loss")),
    "throughput.yaml": ("/app/bi/globalThroughputNetwork", _coverage_kpi("throughput")),
    "network_availability.yaml": ("/app/bi/networkAvailabilityMobile", _coverage_kpi("Network availability")),
    "speed_test.yaml": ("/app/bi/speedTestThruMobile", _coverage_kpi("Speed test")),
    "web_performance.yaml": ("/app/bi/webPerformanceTimesMobile",
                             [(n, {**st, "dimension": "request_time"}) for n, st in _coverage_kpi("Time to first byte")]),
    "video_streaming.yaml": ("/app/bi/youtubeTimesMobile",
                             [(n, {**st, "dimension": "start_time"}) for n, st in _coverage_kpi("Video Start time")]),
    "sample.yaml": ("/app/bi/sample", [("Sample_Net_All", {}), ("Sample_Province_All", {})]),
    "mobile_quality_score.yaml": ("/app/bi/compositeMobileQualityScore",
                                  [(f"{q}_{lvl}_All", {"mqs": q}) for lvl in ("Net", "Province")
                                   for q in ("Excellent", "Sufficient", "Insufficient")]),
    "topology_stock.yaml": ("/app/bi/topologyStock",
                            [(f"Topology Stock_{lvl}_{t}", {"technology": tech, "date": "null"}) for lvl in ("Net", "Province")
                             for t, tech in (("5G", "NR"), ("4G", "LTE"), ("3G", "UMTS"), ("2G", "GSM"))]),
}


@pytest.mark.parametrize("file_name", sorted(KPI_SCENARIOS))
def test_monthly_kpi_scenarios(config, file_name):
    """Every file of scenarios/ from Weplan_export.docx: names and the filters actually applied."""
    import datetime as dt
    path, expected = KPI_SCENARIOS[file_name]
    scenario_file = Path(__file__).parent.parent / "scenarios" / file_name
    markets = {"markets": [{"code": "VTB", "country": "bi", "name": "Burundi"}]}
    results = run_scenarios(load_scenarios([scenario_file], markets), config)
    assert [r.status for r in results] == ["PASS"] * len(expected), [r.error for r in results]

    last = dt.date.today().replace(day=1) - dt.timedelta(days=1)
    out = Path(config["output_dir"]) / "VTB"
    assert len(list(out.glob("*.xlsx"))) == len(expected)
    for name, state in expected:
        f = out / f"VTB_{last.year}_T{last.month}_{name}.xlsx"
        assert f.exists(), f
        data = _xlsx_dict(f)
        assert data["kpi"] == path
        assert data["view"] == ("byCountry" if "_Net_" in name else "byRegions")
        assert data["carrier"] == "ECONET|LUMITEL|ONAMOB|SMART"
        for key, value in state.items():
            assert (data.get(key) or "") == value, (name, key)


def test_scenarios_run_for_each_market(config):
    """Markets from config.yaml: country switched per market, files per market folder, market by market."""
    import datetime as dt
    from weplan_export.config import order_by_market
    markets = [{"code": "VTC", "country": "kh", "name": "Cambodia"},
               {"code": "VTB", "country": "bi", "name": "Burundi"}]
    files = [Path(__file__).parent.parent / "scenarios" / f for f in ("sample.yaml", "coverage_time.yaml")]
    scenarios = order_by_market(load_scenarios(files, {"markets": markets}), markets)
    assert [s.name.split("_")[0] for s in scenarios] == ["VTC"] * 8 + ["VTB"] * 8
    results = run_scenarios(scenarios, config)
    assert all(r.status == "PASS" for r in results), [r.error for r in results]

    last = dt.date.today().replace(day=1) - dt.timedelta(days=1)
    out = Path(config["output_dir"])
    carriers = {"VTC": ("kh", "CELLCARD|METFONE|SMART"), "VTB": ("bi", "ECONET|LUMITEL|ONAMOB|SMART")}
    for code, (country, carrier) in carriers.items():
        got = sorted(f.name for f in (out / code).glob("*.xlsx"))
        assert len(got) == 8, got
        for f in (out / code).glob("*.xlsx"):
            assert f.name.startswith(f"{code}_{last.year}_T{last.month}_")
            data = _xlsx_dict(f)
            assert data["country"] == country, f.name
            assert data["carrier"] == carrier, f.name


def test_empty_table_is_exported_with_warning(tmp_path, config, capsys):
    doc = {"scenarios": [{"name": "haiti_5g", "steps": [
        {"goto": "/app/ht/coverage"},
        {"select_country": "ht"},
        {"set_date": {"preset": "Last month"}},
        {"select_filter": {"id": "coverage_filter", "options": ["5G_SA"]}},
        {"choose_view": "macro"},
        {"wait_for_table": {"empty_after_ms": 3000}},
        {"download_table": {"format": "xlsx", "filename": "empty"}},
    ]}]}
    r = _run(tmp_path, config, doc)[0]
    assert r.ok, r.error
    assert "table is empty" in capsys.readouterr().out
    assert Path(r.downloads[0]["file"]).name == "empty.xlsx"


def test_export_refused_when_country_changed(tmp_path, config):
    """Guard: a page opened under another country's URL must not be exported under this market's name."""
    doc = {"scenarios": [{"name": "wrong_country", "steps": [
        {"goto": "/app/kh/coverage"},
        {"select_country": "kh"},
        {"goto": "/app/bi/coverage"},          # e.g. a hard-coded URL of another country
        {"set_date": {"preset": "Last month"}},
        {"choose_view": "macro"},
        {"wait_for_table": {}},
        {"download_table": {"format": "xlsx"}},
    ]}]}
    r = _run(tmp_path, config, doc)[0]
    assert not r.ok
    assert "shows country 'bi' but this scenario is for 'kh'" in r.error
    assert not r.downloads


def test_resume_after_interruption(tmp_path, server, capsys, monkeypatch):
    """Report written after every scenario; --resume skips what finished and runs the rest."""
    import json
    from weplan_export import actions
    from weplan_export.cli import main

    out = tmp_path / "out"
    cfg = tmp_path / "config.yaml"
    cfg.write_text(yaml.safe_dump({"base_url": server, "output_dir": str(out), "auth": {"required": False}}))
    sc = tmp_path / "s.yaml"
    sc.write_text(yaml.safe_dump({"scenarios": [
        {"name": "a", "steps": [{"goto": "/app/bi/coverage"}]},
        {"name": "b", "steps": [{"goto": "/app/bi/coverage"}, {"check_report": None}, {"interrupt": None}]},
        {"name": "c", "steps": [{"goto": "/app/bi/coverage"}]},
    ]}))

    seen_mid_run = {}

    def check_report(ctx, args):
        report = sorted((out / "_runs").glob("*/report.json"))[-1]
        seen_mid_run.update({r["scenario"]: r["status"] for r in json.loads(report.read_text())})

    def interrupt(ctx, args):
        raise KeyboardInterrupt  # someone stops the run

    monkeypatch.setitem(actions.STEPS, "check_report", check_report)
    monkeypatch.setitem(actions.STEPS, "interrupt", interrupt)
    assert main(["-c", str(cfg), "run", str(sc)]) != 0
    # While b was running, the report already had a finished and b/c pending.
    assert seen_mid_run == {"a": "PASS", "b": "NOT RUN", "c": "NOT RUN"}
    assert "--resume" in capsys.readouterr().out

    monkeypatch.setitem(actions.STEPS, "interrupt", lambda ctx, args: None)
    assert main(["-c", str(cfg), "run", str(sc), "--resume"]) == 0
    text = capsys.readouterr().out
    assert "1 scenario(s) already done, 2 left" in text
    assert "=== [1/2] b" in text and "=== [2/2] c" in text and "=== [1/3] a" not in text
    report = sorted((out / "_runs").glob("*/report.json"))[-1]
    assert {r["scenario"]: r["status"] for r in json.loads(report.read_text())} == \
        {"a": "DONE", "b": "PASS", "c": "PASS"}

    assert main(["-c", str(cfg), "run", str(sc), "--resume"]) == 0
    assert "Nothing left to run." in capsys.readouterr().out


def _market_scenarios(codes):
    return {"scenarios": [{"name": "${market.code}_${kpi}",
                           "matrix": {"market": [{"code": c} for c in codes], "kpi": ["a", "b"]},
                           "steps": [{"goto": "/app/bi/coverage"}]}]}


def test_pause_between_exports_and_after_each_market(tmp_path, config, capsys):
    config["throttle"].update(pause_between=0.2, pause_after_market=[0.6, 0.7])
    t0 = time.time()
    f = tmp_path / "s.yaml"
    f.write_text(yaml.safe_dump(_market_scenarios(["VTC", "STL"]), sort_keys=False), encoding="utf-8")
    results = run_scenarios(load_scenarios([f]), config)  # VTC_a, VTC_b, STL_a, STL_b
    assert [r.status for r in results] == ["PASS"] * 4
    out = capsys.readouterr().out
    assert out.count("Rest between exports: pausing") == 2
    assert out.count("Market VTC done, next STL: pausing") == 1
    assert time.time() - t0 >= 0.2 * 2 + 0.6


def _kill_browser_later(seconds):
    """Kill this test's Chromium processes (children of the Playwright driver) like a user closing it."""
    import os
    import signal
    parents = {}
    for d in Path("/proc").iterdir():
        if d.name.isdigit():
            try:
                parents[int(d.name)] = int((d / "stat").read_text().rsplit(")", 1)[1].split()[1])
            except (OSError, IndexError, ValueError):
                pass

    def descends(pid):
        while pid in parents and pid > 1:
            pid = parents[pid]
            if pid == os.getpid():
                return True
        return False

    victims = [pid for pid in parents if descends(pid) and "chrom" in Path(f"/proc/{pid}/comm").read_text()]
    assert victims

    def kill():
        for pid in victims:
            try:
                os.kill(pid, signal.SIGKILL)
            except OSError:
                pass

    threading.Timer(seconds, kill).start()


@pytest.mark.skipif(not Path("/proc/self/stat").exists(), reason="needs /proc")
def test_closing_the_browser_during_a_pause_stops_the_run(tmp_path, config, monkeypatch):
    from weplan_export import actions
    config["throttle"]["pause_between"] = 30
    monkeypatch.setitem(actions.STEPS, "close_soon", lambda ctx, args: _kill_browser_later(1))
    doc = {"scenarios": [
        {"name": "a", "steps": [{"goto": "/app/bi/coverage"}, {"close_soon": None}]},
        {"name": "b", "steps": [{"goto": "/app/bi/coverage"}]},
    ]}
    t0 = time.time()
    results = _run(tmp_path, config, doc)
    assert [r.status for r in results] == ["PASS", "NOT RUN"]
    assert time.time() - t0 < 20


def test_max_exports_stops_the_run_for_resume(tmp_path, config, capsys):
    config["throttle"]["max_exports"] = 2
    doc = {"scenarios": [{"name": "exp_${n}", "matrix": {"n": [1, 2, 3]}, "steps": [
        {"goto": "/app/bi/coverage"},
        {"set_date": {"preset": "Last month"}},
        {"choose_view": "macro"},
        {"wait_for_table": {}},
        {"download_table": {"format": "csv", "filename": "f_${n}"}},
    ]}]}
    results = _run(tmp_path, config, doc)
    assert [r.status for r in results] == ["PASS", "PASS", "NOT RUN"]
    out = capsys.readouterr().out
    assert "reached throttle.max_exports (2 files)" in out and "--resume" in out


def test_suspended_account_stops_the_whole_run(tmp_path, config, capsys):
    doc = {"scenarios": [
        {"name": "a", "steps": [{"goto": "/app/bi/coverage"}]},
        {"name": "b", "steps": [{"goto": "/suspended"}, {"choose_view": "macro"}]},
        {"name": "c", "steps": [{"goto": "/app/bi/coverage"}]},
    ]}
    results = _run(tmp_path, config, doc)
    assert [r.status for r in results] == ["PASS", "STOPPED", "NOT RUN"]
    assert "temporarily suspended" in results[1].error
    assert results[1].failed_step == "2. choose_view"
    assert results[1].artifacts  # screenshot of the notice
    assert "Stopping the whole run" in capsys.readouterr().out


def test_suspended_at_start_does_not_retry_login(tmp_path, config, monkeypatch):
    from weplan_export import runner
    config["auth"]["required"] = True
    config["start_path"] = "/suspended"
    calls = []
    monkeypatch.setattr(runner, "auto_login", lambda page, cfg: calls.append(1) or False)
    doc = {"scenarios": [{"name": n, "steps": [{"goto": "/app/bi/coverage"}]} for n in "abc"]}
    results = _run(tmp_path, config, doc)
    assert [r.status for r in results] == ["STOPPED", "NOT RUN", "NOT RUN"]
    assert "temporarily suspended" in results[0].error
    assert not calls


def test_not_logged_in_stops_instead_of_retrying(tmp_path, config, monkeypatch):
    from weplan_export import runner
    config["auth"]["required"] = True
    config["selectors"]["logged_in_marker"] = "#never-there"
    calls = []
    monkeypatch.setattr(runner, "auto_login", lambda page, cfg: calls.append(1) or False)
    doc = {"scenarios": [{"name": n, "steps": [{"goto": "/app/bi/coverage"}]} for n in "abc"]}
    results = _run(tmp_path, config, doc)
    assert [r.status for r in results] == ["STOPPED", "NOT RUN", "NOT RUN"]
    assert "Not logged in" in results[0].error
    assert len(calls) == 1


def test_pause_between_steps(tmp_path, config):
    config["throttle"]["pause_between_steps"] = [0.4, 0.5]
    doc = {"scenarios": [{"name": "a", "steps": [
        {"goto": "/app/bi/coverage"}, {"wait": 0}, {"wait": 0}]}]}
    t0 = time.time()
    results = _run(tmp_path, config, doc)
    assert results[0].ok, results[0].error
    assert time.time() - t0 >= 0.8  # 2 pauses: before step 2 and step 3, none before step 1
